from openpyxl import Workbook, load_workbook

from ..config import FILES_DIR
from . import Tool, register


@register
class UpdateTracker(Tool):
    name = "update_tracker"
    description = "Append a row to the event tracker spreadsheet (.xlsx, downloadable). Undo clears the row."
    args_schema = {"sheet": "sheet name", "row": "object of column name -> value"}
    reversible = True

    def run(self, args, ctx):
        FILES_DIR.mkdir(parents=True, exist_ok=True)
        fname = f"tracker_{ctx.get('run_id', 'run')}.xlsx"
        path = FILES_DIR / fname
        sheet = str(args.get("sheet") or "Tracker")[:31]
        row = args.get("row")
        if not isinstance(row, dict):
            row = {k: v for k, v in args.items() if k != "sheet"} or {"note": "update"}
        wb = load_workbook(path) if path.exists() else Workbook()
        if sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            ws = wb.create_sheet(sheet)
            if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
                del wb["Sheet"]
        header = [c.value for c in ws[1]] if ws.max_row >= 1 and ws["A1"].value is not None else []
        for k in row:
            if k not in header:
                header.append(k)
        for i, h in enumerate(header, 1):
            ws.cell(row=1, column=i, value=h)
        idx = max(ws.max_row + 1, 2)
        for k, v in row.items():
            ws.cell(row=idx, column=header.index(k) + 1, value=v if isinstance(v, (int, float, str)) else str(v))
        wb.save(path)
        return {"ok": True, "summary": f"Tracker updated: added a row to sheet '{sheet}'",
                "data": {"file": fname}, "undo_payload": {"file": fname, "sheet": sheet, "row": idx}}

    def undo(self, record):
        u = record["undo_payload"]
        path = FILES_DIR / u["file"]
        if not path.exists():
            return {"ok": False, "summary": "Tracker file not found"}
        wb = load_workbook(path)
        ws = wb[u["sheet"]]
        for c in ws[u["row"]]:
            c.value = None
        wb.save(path)
        return {"ok": True, "summary": "Tracker row cleared"}
